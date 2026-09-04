from pathlib import Path
import json
import sys

from docx import Document


def main() -> None:
    source = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(source)
    payload = {
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else "", "text": p.text}
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip()
        ],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in doc.tables
        ],
        "sections": [
            {
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
            }
            for section in doc.sections
        ],
        "inline_shapes": len(doc.inline_shapes),
    }
    (out_dir / "content.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        suffix = Path(str(part.partname)).suffix or ".bin"
        target = out_dir / f"image-{len(list(out_dir.glob('image-*'))) + 1}{suffix}"
        target.write_bytes(part.blob)


if __name__ == "__main__":
    main()
