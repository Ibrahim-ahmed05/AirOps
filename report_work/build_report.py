from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Administrator\Desktop\Scalability Sprint")
OUT = ROOT / "Scalability_Sprint_Ibrahim_Regenerated.docx"
SHOTS = ROOT / "report_work" / "screenshots"

NAVY = "10233F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
LIGHT_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "17212B"
GREEN = "166534"
PALE_GREEN = "ECFDF3"
AMBER = "7A5A00"
PALE_AMBER = "FFF8E1"
RED = "9B1C1C"
PALE_RED = "FEF2F2"
WHITE = "FFFFFF"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: Sequence[int], indent=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_DXA, (widths_dxa, sum(widths_dxa))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_cell_text(cell, size=8.7, color=DARK, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[int], font_size=8.7):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        shade_cell(cell, PALE_BLUE)
        style_cell_text(cell, size=font_size, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = str(value)
            style_cell_text(cells[idx], size=font_size)
    set_table_geometry(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=DARK)
    set_table_geometry(table, [CONTENT_DXA])
    set_table_borders(table, color="D6E4F0", size="5")
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, text="", bold_prefix=None, italic=False, color=DARK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, color=color, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F7F7F8")
    p_pr.append(shd)
    for line_no, line in enumerate(text.splitlines()):
        if line_no:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.5, color="344054")
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MID_GRAY)


def set_image_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_figure(doc, path: Path, caption: str, width=6.45, alt=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    set_image_alt(shape, alt or caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MID_GRAY, italic=True)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("AirOps | Scalability & Production Readiness")
    set_run_font(hr, size=9, color=MID_GRAY, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Ibrahim | September 3, 2026 | Page ")
    set_run_font(fr, size=9, color=MID_GRAY)
    add_page_number(fp)


def build():
    doc = Document()
    setup_styles(doc)

    # Cover - editorial report pattern, deliberately simple.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SCALABILITY & PRODUCTION\nREADINESS SPRINT")
    set_run_font(r, size=27, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("AirOps - Frontend + Data Performance Report")
    set_run_font(r, size=15, color=BLUE, bold=True)

    add_callout(
        doc,
        "Overall answer",
        "YES WITH ARCHITECTURAL CHANGES. The interface and data-access pattern handle a 520,000-flight dataset, but the current single API process and unoptimized database path are not reliable under concentrated concurrent traffic.",
        fill=PALE_AMBER,
        label_color=AMBER,
    )

    for label, value in [
        ("Member", "Ibrahim - Frontend + Data Performance Lead"),
        ("Project", "AirOps airline operations scalability lab"),
        ("Verified dataset", "520,000 flights; 2,138,472 total rows across five tables"),
        ("Audit date", "September 3, 2026"),
        ("Evidence", "Repository audit, live API tests, PostgreSQL plans, UI checks and a local break probe"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        a = p.add_run(f"{label}: ")
        set_run_font(a, size=10.5, color=MID_GRAY, bold=True)
        b = p.add_run(value)
        set_run_font(b, size=10.5, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(50)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Build -> Measure -> Break -> Find the bottleneck -> Improve -> Measure again")
    set_run_font(r, size=10, color=MID_GRAY, italic=True)
    doc.add_page_break()

    # Executive one-page answer.
    add_heading(doc, "Executive scalability assessment", 1)
    add_callout(
        doc,
        "Decision",
        "The project proves that a bounded frontend can work with a very large dataset. It does not yet prove production capacity at 50-1,000 concurrent users. In the local break test, database pool pressure and expensive list/search queries caused 503 responses and ~10-second tail latency.",
        fill=PALE_RED,
        label_color=RED,
    )
    add_table(
        doc,
        ["Question", "Evidence-based answer"],
        [
            ("What was tested?", "520k data volume, live UI, API timing, payload size, SQL plans, 10/50/100/150-client local read probe, API stop/restart."),
            ("What worked?", "Pagination keeps the DOM bounded; dashboard/list/detail pages work; current payload is small; errors are visible; service recovered after load and restart."),
            ("What broke?", "At 50+ concentrated clients, requests queued behind the 20-connection pool and many returned HTTP 503 after connection timeouts."),
            ("Biggest bottleneck", "PostgreSQL work per request: full scans/sorts, wildcard ILIKE search, exact COUNT on each page, on-demand dashboard aggregates and limited pool capacity."),
            ("What was improved?", "Frontend caching/deduplication, stable keys, 300 ms search debounce, session airport cache, smaller list projection, memoization and cache invalidation."),
            ("Current comfortable capacity", "Not certified. A 10-client probe had 0% errors but p95 was ~5.0 s because search was slow; 50 clients were unstable in this environment."),
            ("Next scaling step", "Fix the load scripts, add evidence-led indexes and query changes, reduce count/aggregate work, add backpressure/monitoring, then rerun 50/100/150 before attempting 500/1,000."),
        ],
        [2450, 6910],
        font_size=8.5,
    )

    add_heading(doc, "Capacity answer", 2)
    add_table(
        doc,
        ["Load", "Current status", "What it means"],
        [
            ("10", "Degraded", "0% errors, but p95 5.01 s in the short local mixed read probe."),
            ("50", "Unstable", "p95 10.08 s; 50.0% failures/timeouts in the same probe."),
            ("100", "Unstable", "p95 10.38 s; 92.35% failures."),
            ("150", "Unstable", "p95 10.66 s; 91.58% failures."),
            ("500", "Not run", "Requires query/index work, cached aggregates, pooler, load shedding and horizontally scaled API."),
            ("1,000", "Not run", "Requires the 500-user changes plus production-scale validation, autoscaling, read scaling and stronger observability."),
        ],
        [1000, 1650, 6710],
        font_size=8.2,
    )
    add_para(doc, "Important: these are local, 10-second read-only probe results with 750 ms think time. They expose a real failure mode but are not a replacement for repeatable k6 tests on production-like infrastructure.", italic=True, color=MID_GRAY, after=2)
    doc.add_page_break()

    # Introduction and architecture.
    add_heading(doc, "1. Project introduction", 1)
    add_para(doc, "AirOps is a deliberately small but production-style airline operations application. The goal is not feature volume; the goal is to watch user experience and system reliability as the dataset and request concurrency increase.")
    add_para(doc, "Ibrahim's part of the sprint asks a practical question: what does the user see when the backend is slow, overloaded or unavailable? The work therefore covers the large-data interface, API behavior from the browser, failure states, request reduction and evidence for where the architecture stops being reliable.")

    add_heading(doc, "1.1 Architecture used", 2)
    add_table(
        doc,
        ["Layer", "Technology", "Purpose"],
        [
            ("Frontend", "Next.js 14, React 18, TypeScript, Tailwind", "Dashboard, bounded table, filters, detail/update flow, diagnostics and failure UX."),
            ("Client data", "TanStack Query", "Caching, in-flight deduplication, background refresh, retries and invalidation."),
            ("API", "Node.js + Fastify", "Validated REST endpoints, structured logs, CORS, timing hooks and graceful shutdown."),
            ("Database", "PostgreSQL 16 + Drizzle", "Relational flight/airport/aircraft/event/incident data; pool max is 20."),
            ("Test tools", "Benchmark script, EXPLAIN ANALYZE, load scripts", "Single-user baseline, query evidence and planned concurrent-user scenarios."),
        ],
        [1400, 2800, 5160],
    )

    add_heading(doc, "1.2 Verified large dataset", 2)
    add_table(
        doc,
        ["Table", "Verified rows", "Notes"],
        [
            ("flights", "520,000", "Primary operations dataset."),
            ("flight_events", "1,602,332", "Historical status/activity records."),
            ("incidents", "15,850", "Operational incident records."),
            ("airports", "40", "Static reference data."),
            ("aircraft", "250", "Fleet reference data."),
            ("Total", "2,138,472", "All five relational tables."),
        ],
        [1900, 1700, 5760],
    )
    add_callout(doc, "Documentation mismatch", "The seed code currently inserts 2,000 flights per batch. The older data-model document says 5,000. The report uses the code as the source of truth.", fill=PALE_AMBER, label_color=AMBER)
    doc.add_page_break()

    # What was built.
    add_heading(doc, "2. What I built", 1)
    add_heading(doc, "2.1 High-volume frontend", 2)
    for item in [
        "Operations dashboard with eight KPI cards, status distribution and manual/background refresh.",
        "Flight table with server-side pagination (20 by default, API maximum 100), search, status/origin/destination filters and sorting.",
        "Flight detail view with airport, aircraft, timings, events, incidents and a status-update workflow.",
        "Diagnostics page showing API reachability, database state, query latency and process uptime.",
        "The browser never tries to render 520,000 rows. Only the current page is fetched and rendered.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.2 API and data foundation", 2)
    for item in [
        "GET /health, /api/airports, /api/flights, /api/flights/:id and /api/dashboard, plus PATCH /api/flights/:id/status.",
        "Zod validation, centralized JSON errors, Pino request logging, request-duration logging, CORS and graceful shutdown.",
        "Five-table Drizzle schema with foreign keys and cascading event/incident cleanup.",
        "Bulk seeder designed for 10k, 100k, 500k and 1m flights, with child events/incidents generated in chunks.",
        "Status update and event insertion are wrapped in one database transaction, which protects against partial writes.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.3 Frontend performance and resilience work", 2)
    for item in [
        "TanStack Query cache and in-flight request deduplication.",
        "Stable primitive query keys so a real filter/page change triggers one new request.",
        "300 ms search debounce to avoid one request per keystroke.",
        "Airports cached for the browser session (staleTime and gcTime set to Infinity on the page query).",
        "Flight-list response projected to only fields rendered by the table.",
        "Memoized metric cards and flight rows; targeted cache invalidation after a status change.",
        "Loading skeletons, retry actions, empty states, timeout/cancellation and development-only failure injection.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "3. Screenshots from the live 520k dataset", 1)
    add_figure(doc, SHOTS / "dashboard.png", "Figure 1 - Dashboard reading real metrics from 520,000 flights.", 6.45, "AirOps dashboard showing 520,000 total flights and operational KPIs")
    add_para(doc, "The dashboard uses one compact API response (0.3 KB in the benchmark), but the database still recomputes full-table aggregates on each refresh.", italic=True, color=MID_GRAY)
    doc.add_page_break()

    add_heading(doc, "3.1 Bounded flight table", 1)
    add_figure(doc, SHOTS / "flights.png", "Figure 2 - Paginated table: 20 visible rows out of 520,000.", 5.25, "AirOps flight operations table showing pagination, filters and sorting")
    add_callout(doc, "Why no table virtualization", "Pagination already caps each response and DOM render to at most 100 rows. Virtualization would add complexity without fixing the real bottleneck, which is server/database query work.")
    doc.add_page_break()

    # Testing methodology.
    add_heading(doc, "4. What I tested", 1)
    add_table(
        doc,
        ["Check", "Result", "Evidence"],
        [
            ("API TypeScript", "PASS", "tsc --noEmit completed."),
            ("Frontend TypeScript", "PASS", "tsc --noEmit completed."),
            ("API production build", "PASS", "tsup built dist/server.js in 82 ms."),
            ("Frontend production build", "PASS", "Next.js generated all routes successfully."),
            ("Live database", "PASS", "Health endpoint connected; 520,000 flights returned by pagination."),
            ("Single-user API baseline", "PASS", "10 warm-run samples per endpoint."),
            ("SQL execution plans", "PASS", "EXPLAIN ANALYZE captured on the live 520k database."),
            ("Concurrent break probe", "PARTIAL", "10/50/100/150 read-only clients; short local run, not formal k6."),
            ("Service restart", "PASS", "UI showed offline state and returned online after restart."),
            ("k6 suite", "BLOCKED", "Saved smoke run made 0 requests because request configuration is invalid."),
        ],
        [2300, 1200, 5860],
        font_size=8.4,
    )

    add_heading(doc, "4.1 Production frontend output", 2)
    add_table(
        doc,
        ["Route", "Route JS", "First-load JS"],
        [
            ("/dashboard", "0.18 KB", "116 KB"),
            ("/flights", "4.32 KB", "108 KB"),
            ("/flights/[id]", "3.62 KB", "116 KB"),
            ("/diagnostics", "3.57 KB", "90.8 KB"),
            ("Shared", "-", "87.2 KB"),
        ],
        [3300, 2300, 3760],
    )
    add_para(doc, "These build sizes confirm route-level splitting is working. They do not replace real browser FCP/LCP/INP measurements, which are still missing.", italic=True, color=MID_GRAY)
    doc.add_page_break()

    # Baseline actual measurements.
    add_heading(doc, "5. Actual 520k-record API baseline", 1)
    add_para(doc, "The following results were measured locally on September 3, 2026 after three warm-up requests. Each endpoint was sampled 10 times. These are current-state timings, not before/after optimization results.")
    add_table(
        doc,
        ["Endpoint", "Avg", "P95", "Payload", "Result"],
        [
            ("GET /health", "2.1 ms", "3.2 ms", "0.1 KB", "Fast"),
            ("GET /api/airports", "2.5 ms", "2.9 ms", "4.6 KB", "Fast; cached in browser"),
            ("GET /api/dashboard", "98.8 ms", "109.5 ms", "0.3 KB", "Good alone; full scan"),
            ("Flights: 20 rows", "273.6 ms", "290.1 ms", "6.7 KB", "Usable alone; full sort"),
            ("Flights: 50 rows", "346.7 ms", "450.8 ms", "16.6 KB", "Higher payload + same scan"),
            ("Page 50 / 20 rows", "290.0 ms", "332.4 ms", "6.7 KB", "Offset cost starting"),
            ("Status = DELAYED", "126.6 ms", "139.2 ms", "6.6 KB", "Sequential status scan"),
            ("Search = AA", "951.9 ms", "990.7 ms", "6.7 KB", "Slowest read path"),
            ("Flight detail", "86.7 ms", "151.9 ms", "0.9 KB", "PK lookup is efficient"),
        ],
        [3550, 1250, 1250, 1350, 1960],
        font_size=8.1,
    )
    add_callout(doc, "User-experience meaning", "The table feels reasonable for a single user, but a search already consumes about one second before frontend rendering. Under concurrency, these expensive searches occupy pool connections and amplify queueing.", fill=PALE_AMBER, label_color=AMBER)

    add_heading(doc, "5.1 Payload result", 2)
    add_para(doc, "The current 20-row flight-list response is 6,826 bytes (6.7 KB). This is better than the older documentation's 10-22 KB estimates. A true before/after percentage cannot be claimed because the pre-projection response was not saved or rerun.")
    doc.add_page_break()

    # Query plans.
    add_heading(doc, "6. PostgreSQL bottleneck evidence", 1)
    add_para(doc, "EXPLAIN ANALYZE confirms that the main performance cost sits below the React interface. The database repeatedly scans and sorts the 520,000-row flight table.")
    add_table(
        doc,
        ["Query", "Plan evidence", "Execution"],
        [
            ("Dashboard aggregates", "Parallel sequential scan of all flights", "101.7 ms"),
            ("Exact pagination COUNT", "Parallel sequential scan of all flights", "42.6 ms"),
            ("Page 1 sorted list", "Full parallel scan + 3 joins + top-N sort", "292.5 ms"),
            ("Deep offset (980)", "Full scan/sort and discard earlier rows", "244.7 ms"),
            ("ILIKE %AA% search", "Full scan + airport joins; 158,920 rows/worker rejected", "429.9 ms"),
            ("Status count", "Parallel sequential scan; no status index", "47.1 ms"),
            ("Detail by ID", "Primary-key index + small joins", "3.0 ms"),
        ],
        [2700, 4660, 2000],
        font_size=8.5,
    )
    add_heading(doc, "Why this becomes a cliff", 2)
    for item in [
        "Every list request performs two pieces of work: an exact COUNT and the data query.",
        "Sorting by scheduled departure has no supporting index, so returning 20 rows can still scan/sort the full table.",
        "Leading-wildcard ILIKE search cannot use a normal B-tree index and was the slowest endpoint.",
        "Dashboard refreshes perform live aggregates instead of reading cached or precomputed metrics.",
        "The API pool allows 20 database connections. Expensive concurrent requests fill it; later requests wait and time out after 5 seconds.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    # Break test.
    add_heading(doc, "7. Break-it test: concentrated mixed reads", 1)
    add_para(doc, "I ran a short local read-only probe against the 520k database using 60% paginated flight reads, 20% wildcard searches, 10% dashboard requests and 10% health checks. Each client waited 750 ms between requests. Requests had a 10-second timeout.")
    add_table(
        doc,
        ["Clients", "Requests", "P50", "P95", "P99", "Failures"],
        [
            ("10", "42", "1.22 s", "5.01 s", "7.93 s", "0.00%"),
            ("50", "92", "7.51 s", "10.08 s", "10.08 s", "50.00%"),
            ("100", "170", "6.77 s", "10.38 s", "10.40 s", "92.35%"),
            ("150", "273", "5.81 s", "10.66 s", "10.67 s", "91.58%"),
        ],
        [1100, 1500, 1500, 1500, 1500, 2260],
        font_size=8.8,
    )
    add_callout(doc, "What actually broke", "API logs repeatedly reported 'timeout exceeded when trying to connect'. The server returned HTTP 503 after roughly 5.25 seconds while the database pool was saturated. Search failed first and most consistently.", fill=PALE_RED, label_color=RED)
    add_heading(doc, "Recovery", 2)
    add_para(doc, "After the probe stopped, the API returned 200 responses again without being restarted. The health payload reported the database connected with a 1.56 ms ping. This shows the failure was capacity/queue pressure rather than permanent corruption or a process crash.")
    add_heading(doc, "Limits of this result", 2)
    for item in [
        "It is a local probe on one machine, not a production-sized distributed load test.",
        "The run is short, so it does not prove memory stability or sustained capacity.",
        "It is read-only and does not prove write safety or lock behavior.",
        "It is still useful because it reproduced a specific, logged failure mechanism and recovery path.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    # Service restart UI screenshots.
    add_heading(doc, "8. Failure handling seen by the user", 1)
    add_para(doc, "I stopped the API while the production frontend remained open. Within the diagnostics polling cycle, the interface changed to API OFFLINE / DB DISCONNECTED, showed 'Failed to fetch', kept navigation working and exposed a Check Now action.")
    add_figure(doc, SHOTS / "diagnostics-offline.png", "Figure 3 - API stopped: clear offline state instead of a broken page.", 6.25, "AirOps diagnostics page showing API offline and database disconnected")
    add_para(doc, "I then restarted the API. The next polling cycle changed the page back to ONLINE / CONNECTED and displayed fresh uptime and database latency without reloading the frontend.")
    add_figure(doc, SHOTS / "diagnostics-recovered.png", "Figure 4 - API restarted: automatic recovery detected by the UI.", 6.25, "AirOps diagnostics page showing API online after service restart")
    doc.add_page_break()

    add_heading(doc, "8.1 Implemented failure behaviors", 1)
    add_table(
        doc,
        ["Condition", "Current behavior", "Evidence status"],
        [
            ("Loading / slow response", "Skeleton rows/cards and loading indicators.", "Implemented in shared components and pages."),
            ("HTTP 500 / 503", "Typed error message, warning state and retry action.", "Implemented; 503 reproduced during break probe."),
            ("Network/API down", "Offline diagnostics and network-specific error text.", "Live stop/restart verified."),
            ("Timeout", "AbortController with 5 s dashboard and 10 s list/detail timeouts.", "Implemented; connection timeout observed server-side."),
            ("Retry", "TanStack Query retries GET failures once; mutations do not retry.", "Code verified. Older docs incorrectly say 3 retries."),
            ("Empty results", "Contextual empty state and filter-reset action.", "Implemented; not re-measured under load."),
            ("Stale response", "Query cancellation and stable keys reduce stale overwrites.", "Code verified."),
            ("Injected failures", "Development-only latency/error/timeout wrapper.", "Implemented; production build intentionally disables it."),
        ],
        [1900, 4400, 3060],
        font_size=8.2,
    )
    add_heading(doc, "Data safety under concurrent writes", 2)
    add_para(doc, "The PATCH route writes the new status and its flight event inside one transaction, so those two operations succeed or fail together. However, the row has no version field or optimistic-lock check. Two writers can overwrite each other (last write wins) while both events remain. The supplied write-contention scenario has not produced valid recorded results, so data safety under competing updates is not yet proven.")
    doc.add_page_break()

    # Optimizations and evidence quality.
    add_heading(doc, "9. Optimizations completed", 1)
    add_table(
        doc,
        ["Change", "Why it helps", "What is proven"],
        [
            ("Server-side pagination", "Bounds network and DOM work.", "20 rows displayed from 520k; max API page is 100."),
            ("300 ms search debounce", "Avoids a request for every keypress.", "Code behavior is clear; no saved browser trace."),
            ("TanStack Query", "Deduplicates/caches GET requests.", "Configuration and page usage verified."),
            ("Airports session cache", "Stops repeated static-data fetches.", "staleTime/gcTime Infinity verified."),
            ("Stable query keys", "Only meaningful state changes refetch.", "Primitive key construction verified."),
            ("List projection", "Cuts unused JSON fields.", "Current 20-row payload measured at 6.7 KB."),
            ("Memoized rows/cards", "Reduces repeat React work.", "Implementation verified; CPU saving not measured."),
            ("Targeted invalidation", "Keeps detail/list data fresh after write.", "Mutation invalidation verified."),
        ],
        [2550, 3400, 3410],
        font_size=8.3,
    )

    add_heading(doc, "9.1 Before/after evidence status", 2)
    add_callout(doc, "Honest result", "A numeric before/after performance claim is still incomplete. The repository contains current measurements and several estimates, but it does not contain a saved pre-optimization build, raw DevTools trace or valid k6 output for comparison.", fill=PALE_AMBER, label_color=AMBER)
    add_table(
        doc,
        ["Metric", "Current / defensible statement", "Do not claim yet"],
        [
            ("Search calls", "Debounce turns a fast 5-keystroke input into about 1 request after idle.", "Measured 80% network reduction without a trace."),
            ("Airport calls", "Page query caches airport data for the session.", "Measured 5-to-1 reduction without recorded navigation."),
            ("Flight payload", "Current 20-row response is 6.7 KB.", "50% reduction unless old payload is reproduced."),
            ("Page speed", "Production build sizes are known.", "25-50% faster FCP/CPU; prior values are estimates."),
            ("Capacity", "The local probe shows a cliff at 50 concentrated clients.", "Stable at 50 based on the existing documentation."),
        ],
        [2100, 4020, 3240],
        font_size=8.2,
    )
    doc.add_page_break()

    # K6 audit.
    add_heading(doc, "10. k6 framework: built, but not yet valid", 1)
    add_para(doc, "The repository contains smoke, mixed-load, progressive-stress, spike and write-contention scenarios with presets from 1 to 1,000 VUs. That is useful scaffolding, but it is not the same as executed evidence.")
    add_table(
        doc,
        ["Issue", "Impact", "Required fix"],
        [
            ("responseType: 'json' in HTTP options", "k6 rejects it; saved smoke run sent 0 requests.", "Remove it or use a supported response type; parse JSON from the response."),
            ("options references options.vus while building itself", "load.js and write-contention.js can fail during module evaluation.", "Calculate targetVUs before exporting options."),
            ("No raw result files", "50/100/150/250 claims cannot be verified.", "Export JSON/CSV and retain console summaries for every tier."),
            ("Threshold rate >100 req/s", "May conflict with deliberate user think time and create false failures.", "Set arrival-rate goals from a workload model or remove the arbitrary threshold."),
            ("k6 unavailable in this environment", "Formal suite could not be rerun here.", "Install/restore k6, then run smoke before higher tiers."),
        ],
        [2450, 3570, 3340],
        font_size=8.3,
    )
    add_para(doc, "The existing smoke-test-output.txt confirms the first issue: setup ended with zero HTTP requests and 'json does not belong to ResponseType values'. Therefore the older 50/100/250 stability table must be treated as expected/example data, not measured performance.")

    add_heading(doc, "Correct validation order", 2)
    for item in [
        "Fix the k6 request options and self-referencing VU configuration.",
        "Run smoke and confirm non-zero requests and 100% expected checks.",
        "Run 10, 50, 100 and 150 VUs on the same build/data/hardware; repeat each important level at least three times.",
        "Capture API logs, PostgreSQL active connections, pool wait/timeout counts, CPU, memory, disk I/O and event-loop lag during the same window.",
        "Apply one optimization batch, then rerun the identical suite for a real before/after comparison.",
        "Only after 150 is repeatably stable should 500 and 1,000-user tests be attempted.",
    ]:
        add_number(doc, item)
    doc.add_page_break()

    # Recommendations.
    add_heading(doc, "11. What must change to scale", 1)
    add_heading(doc, "First: make 50-150 users reliable", 2)
    for item in [
        "Add indexes based on the captured plans: scheduled_departure for sort/pagination, selective status/date/airport combinations for real access patterns, and pg_trgm or a better search path for leading-wildcard search.",
        "Remove unnecessary exact COUNT work from every page where possible; use cached/approximate totals or return hasNext instead.",
        "Move dashboard aggregates to a short-TTL cache, rollup table or materialized view instead of rescanning 520k flights every 10 seconds per client.",
        "Use keyset/cursor pagination for deep navigation instead of OFFSET.",
        "Add request concurrency limits, backpressure and sensible 429/503 responses before the pool is exhausted.",
        "Tune the pool with PostgreSQL limits and measured query duration; add PgBouncer when multiple API processes share the database.",
        "Fix k6 and define an SLO such as p95 <1 s and errors <1% for the critical list/detail paths.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "At about 500 concurrent users", 2)
    for item in [
        "Run multiple stateless Fastify instances behind a load balancer.",
        "Use PgBouncer/connection governance so API scale-out does not create a database connection storm.",
        "Cache hot reference data and dashboard/read-heavy responses; invalidate deliberately after writes.",
        "Add observability for endpoint p95/p99, errors, pool wait time, slow queries, CPU/memory and recovery time.",
        "Protect writes with optimistic concurrency/version checks and idempotency where repeated client actions are possible.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "At about 1,000 concurrent users", 2)
    for item in [
        "Keep horizontally scaled API instances with autoscaling based on latency/CPU/pool pressure, not request count alone.",
        "Serve static frontend assets through CDN/edge caching.",
        "Separate read/reporting pressure from transactional writes; introduce read replicas only after query/index tuning.",
        "Use queues for genuinely asynchronous work, not for normal read requests.",
        "Consider partitioning or specialized search/analytics only if measured plans show that simpler optimizations are no longer enough.",
        "Run sustained, spike, slow-database, restart and write-contention release gates on production-like infrastructure.",
    ]:
        add_bullet(doc, item)
    # Final answer.
    add_heading(doc, "12. Final answer: CAN WE SCALE?", 1)
    add_callout(doc, "Verdict", "YES WITH ARCHITECTURAL CHANGES - and NOT YET a clean production-capacity YES for the current baseline.", fill=PALE_AMBER, label_color=AMBER)
    add_heading(doc, "What the sprint has proved", 2)
    for item in [
        "A well-bounded interface can browse and operate on a 520,000-flight / 2.14-million-row relational dataset.",
        "Frontend request amplification and payload bloat have been reduced in the implementation.",
        "The UI stays understandable when the API is stopped and automatically recognizes recovery.",
        "The main bottleneck is measurable and explainable: expensive database work fills a 20-connection pool, causing queueing and 503 timeouts.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "What is still required for a clean YES", 2)
    for item in [
        "A working k6 suite with saved, repeatable results at 50, 100 and 150 users.",
        "A true pre-change benchmark and matching post-change benchmark.",
        "Indexes/query changes/cached aggregates proven with EXPLAIN ANALYZE and load results.",
        "Write-contention evidence proving no lost updates or inconsistent history.",
        "Monitoring evidence for failure detection and measured recovery time.",
        "Only then: controlled 500 and 1,000-user tests on production-like infrastructure.",
    ]:
        add_bullet(doc, item)

    add_para(doc, "Vision71 can take substantially larger applications if this method becomes the standard: measure honestly, expose the bottleneck, fix the smallest effective layer, and repeat the same test. The current project shows the right direction, but the final capacity claim must be earned with repeatable evidence.")

    add_heading(doc, "Evidence reviewed", 2)
    add_para(doc, "Repository source and Markdown documentation; live API responses; apps/api/scripts/benchmark.mjs; apps/api/scripts/query-plans.mjs; load-tests/*.js; smoke-test-output.txt; production builds; and browser screenshots captured from the running application.", italic=True, color=MID_GRAY)

    core = doc.core_properties
    core.title = "AirOps Scalability & Production Readiness Sprint - Ibrahim"
    core.subject = "Frontend, data performance, concurrent load, failure handling and scaling assessment"
    core.author = "Ibrahim"
    core.keywords = "AirOps, scalability, frontend performance, PostgreSQL, Fastify, Next.js"
    core.comments = "Regenerated from verified repository and runtime evidence on September 3, 2026."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
